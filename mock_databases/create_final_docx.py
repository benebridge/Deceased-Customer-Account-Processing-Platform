#!/usr/bin/env python3
"""
Create comprehensive Word document from CSV with all required fields
"""

import csv
import sqlite3
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_reference_docx():
    base_dir = Path(__file__).parent

    # Read CSV with all person data
    csv_file = base_dir / "MASTER_REFERENCE_ALL_PERSONS.csv"
    persons = []
    with open(csv_file, 'r') as f:
        reader = csv.DictReader(f)
        for row in reader:
            persons.append(row)

    # Separate deceased and beneficiaries
    deceased_list = [p for p in persons if p['type'] == 'deceased']
    beneficiaries_list = [p for p in persons if p['type'] == 'beneficiary']

    # Get identity records and account designations from databases
    conn_id = sqlite3.connect(base_dir / "identity_verification.db")
    conn_id.row_factory = sqlite3.Row
    c_id = conn_id.cursor()
    c_id.execute("SELECT * FROM identity_records")
    identities = {dict(row)['ssn']: dict(row) for row in c_id.fetchall()}
    conn_id.close()

    conn_ben = sqlite3.connect(base_dir / "beneficiary_registry.db")
    conn_ben.row_factory = sqlite3.Row
    c_ben = conn_ben.cursor()
    c_ben.execute("SELECT * FROM beneficiary_designations")
    designations_list = [dict(row) for row in c_ben.fetchall()]
    conn_ben.close()

    # Create beneficiary lookup by SSN
    ben_by_ssn = {}
    for b in beneficiaries_list:
        if b['ssn'] not in ben_by_ssn:
            ben_by_ssn[b['ssn']] = []
        ben_by_ssn[b['ssn']].append(b)

    # Get accounts
    conn_fin = sqlite3.connect(base_dir / "financial_accounts.db")
    conn_fin.row_factory = sqlite3.Row
    c_fin = conn_fin.cursor()
    c_fin.execute("SELECT * FROM accounts")
    accounts_list = [dict(row) for row in c_fin.fetchall()]
    conn_fin.close()

    # Create document
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    # Title
    title = doc.add_heading('BeneBridge POC - Person Reference for Forms', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    subtitle.add_run('LA County Death Certificates | CA Driver\'s Licenses | IRA Beneficiary Claim Forms').bold = True
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    summary = doc.add_paragraph()
    summary.add_run(f'Total Cases: 20 | Deceased: {len(deceased_list)} | Beneficiaries: {len(beneficiaries_list)}\n').bold = True
    summary.add_run('All from California | Institution: Community National Bank | Account Type: IRA Only')

    doc.add_page_break()

    # ========== SECTION 1: DECEASED ==========
    doc.add_heading('SECTION 1: DECEASED PERSONS', 1)
    doc.add_paragraph('For LA COUNTY DEATH CERTIFICATES')
    doc.add_paragraph()

    for person in deceased_list:
        heading = doc.add_heading(f'Case #{person["case_id"]}: {person["full_name"]}', 2)
        heading.runs[0].font.color.rgb = RGBColor(0, 0, 128)

        table = doc.add_table(rows=13, cols=2)
        table.style = 'Light Grid Accent 1'
        rows = table.rows

        rows[0].cells[0].text = 'Full Name'
        rows[0].cells[1].text = person['full_name']

        rows[1].cells[0].text = 'First Name'
        rows[1].cells[1].text = person['first_name']

        rows[2].cells[0].text = 'Last Name'
        rows[2].cells[1].text = person['last_name']

        rows[3].cells[0].text = 'Gender'
        rows[3].cells[1].text = 'Male' if person['gender'] == 'M' else 'Female'

        rows[4].cells[0].text = 'SSN'
        rows[4].cells[1].text = person['ssn']

        rows[5].cells[0].text = 'Date of Birth'
        rows[5].cells[1].text = person['dob']

        rows[6].cells[0].text = 'Date of Death'
        rows[6].cells[1].text = person['dod']

        rows[7].cells[0].text = 'Address'
        rows[7].cells[1].text = f"{person['address_street']}, {person['address_city']}, {person['address_state']} {person['address_zip']}"

        rows[8].cells[0].text = 'City/County'
        rows[8].cells[1].text = f"{person['address_city']}, Los Angeles County, CA"

        rows[9].cells[0].text = 'Height'
        rows[9].cells[1].text = person['height']

        rows[10].cells[0].text = 'Weight'
        rows[10].cells[1].text = f"{person['weight']} lbs"

        rows[11].cells[0].text = 'Eye Color'
        rows[11].cells[1].text = person['eye_color']

        rows[12].cells[0].text = 'Sex'
        rows[12].cells[1].text = person['gender']

        # Get accounts for this case
        case_accounts = [a for a in accounts_list if a['deceased_case_id'] == int(person['case_id'])]
        if case_accounts:
            doc.add_paragraph()
            doc.add_paragraph('IRA Accounts:').bold = True
            for acc in sorted(case_accounts, key=lambda x: x['balance'], reverse=True):
                acc_para = doc.add_paragraph(style='List Bullet')
                acc_para.add_run(f"Account #: {acc['account_id']} | Balance: ${acc['balance']:,.2f}")

        doc.add_paragraph()
        doc.add_paragraph('_' * 80)
        doc.add_paragraph()

    doc.add_page_break()

    # ========== SECTION 2: BENEFICIARIES ==========
    doc.add_heading('SECTION 2: BENEFICIARIES', 1)
    doc.add_paragraph('For CA DRIVER\'S LICENSES and IRA BENEFICIARY CLAIM FORMS')
    doc.add_paragraph()

    for person in beneficiaries_list:
        heading = doc.add_heading(f'Beneficiary: {person["full_name"]}', 2)
        heading.runs[0].font.color.rgb = RGBColor(0, 100, 0)

        # Get deceased name for this beneficiary
        deceased = next((d for d in deceased_list if d['case_id'] == person['case_id']), None)
        deceased_name = deceased['full_name'] if deceased else 'Unknown'

        doc.add_paragraph(f"For Case #{person['case_id']}: {deceased_name} | Relationship: {person['relationship']}").italic = True

        # Get identity record
        identity = identities.get(person['ssn'])

        doc.add_paragraph()
        doc.add_paragraph('FOR CA DRIVER\'S LICENSE:').bold = True

        # DL table
        table_dl = doc.add_table(rows=13, cols=2)
        table_dl.style = 'Light Grid Accent 1'
        rows_dl = table_dl.rows

        rows_dl[0].cells[0].text = 'Driver\'s License Number'
        rows_dl[0].cells[1].text = person.get('drivers_license', identity['drivers_license'] if identity else '')

        rows_dl[1].cells[0].text = 'Issue Date'
        rows_dl[1].cells[1].text = identity['dl_issue_date'] if identity else ''

        rows_dl[2].cells[0].text = 'Expiration Date'
        rows_dl[2].cells[1].text = identity['dl_expiration'] if identity else ''

        rows_dl[3].cells[0].text = 'Last Name'
        rows_dl[3].cells[1].text = person['last_name']

        rows_dl[4].cells[0].text = 'First Name'
        rows_dl[4].cells[1].text = person['first_name']

        rows_dl[5].cells[0].text = 'Address'
        rows_dl[5].cells[1].text = f"{person['address_street']}, {person['address_city']}, {person['address_state']} {person['address_zip']}"

        rows_dl[6].cells[0].text = 'Date of Birth'
        rows_dl[6].cells[1].text = person['dob']

        rows_dl[7].cells[0].text = 'Sex'
        rows_dl[7].cells[1].text = person['gender']

        rows_dl[8].cells[0].text = 'Height'
        rows_dl[8].cells[1].text = person.get('height', identity['height'] if identity else '')

        rows_dl[9].cells[0].text = 'Weight'
        weight_val = person.get('weight', identity['weight'] if identity else '')
        rows_dl[9].cells[1].text = f"{weight_val} lbs" if weight_val else ''

        rows_dl[10].cells[0].text = 'Eye Color'
        rows_dl[10].cells[1].text = person.get('eye_color', identity['eye_color'] if identity else '')

        rows_dl[11].cells[0].text = 'Hair Color'
        rows_dl[11].cells[1].text = identity['hair_color'] if identity else ''

        rows_dl[12].cells[0].text = 'Photo'
        rows_dl[12].cells[1].text = identity['photo_reference'] if identity else ''

        doc.add_paragraph()
        doc.add_paragraph('FOR IRA BENEFICIARY CLAIM FORM:').bold = True

        # Claim form table
        table_claim = doc.add_table(rows=7, cols=2)
        table_claim.style = 'Light Grid Accent 1'
        rows_claim = table_claim.rows

        rows_claim[0].cells[0].text = 'Full Name'
        rows_claim[0].cells[1].text = person['full_name']

        rows_claim[1].cells[0].text = 'SSN'
        rows_claim[1].cells[1].text = person['ssn']

        rows_claim[2].cells[0].text = 'Email'
        rows_claim[2].cells[1].text = person.get('email', '')

        rows_claim[3].cells[0].text = 'Phone'
        rows_claim[3].cells[1].text = person.get('phone', '')

        rows_claim[4].cells[0].text = 'Relationship to Deceased'
        rows_claim[4].cells[1].text = person['relationship']

        # Find accounts for this beneficiary
        # Need to find beneficiary_id from beneficiaries table by SSN
        conn_ben2 = sqlite3.connect(base_dir / "beneficiary_registry.db")
        conn_ben2.row_factory = sqlite3.Row
        c_ben2 = conn_ben2.cursor()
        c_ben2.execute("SELECT beneficiary_id FROM beneficiaries WHERE ssn = ?", (person['ssn'],))
        ben_row = c_ben2.fetchone()
        conn_ben2.close()

        if ben_row:
            ben_id = ben_row['beneficiary_id']
            ben_desigs = [d for d in designations_list if d['beneficiary_id'] == ben_id]

            if ben_desigs:
                account_list = [f"{d['account_id']} ({d['percentage']}%)" for d in ben_desigs]
                rows_claim[5].cells[0].text = 'IRA Account Number(s)'
                rows_claim[5].cells[1].text = ', '.join(account_list)

                percentages = ', '.join([f"{d['percentage']}%" for d in ben_desigs])
                rows_claim[6].cells[0].text = 'Beneficiary Percentage(s)'
                rows_claim[6].cells[1].text = percentages

        doc.add_paragraph()
        doc.add_paragraph('_' * 80)
        doc.add_paragraph()

    # Save
    output_file = base_dir / "PERSON_REFERENCE_FOR_FORMS.docx"
    doc.save(output_file)

    print("=" * 80)
    print("WORD DOCUMENT CREATED")
    print("=" * 80)
    print(f"File: {output_file}")
    print(f"Size: {output_file.stat().st_size / 1024:.1f} KB")
    print()
    print("Document contains:")
    print(f"  - {len(deceased_list)} deceased persons")
    print(f"  - {len(beneficiaries_list)} beneficiaries")
    print()
    print("SECTION 1: Deceased with all physical characteristics (no hair color)")
    print("SECTION 2: Beneficiaries with:")
    print("  - CA Driver's License info (DL#, issue/exp dates, hair color, all fields)")
    print("  - IRA Beneficiary Claim Form info (SSN, email, accounts, %)")
    print("  - All beneficiaries from California!")
    print()
    print("Ready for form filling!")

if __name__ == "__main__":
    create_reference_docx()
