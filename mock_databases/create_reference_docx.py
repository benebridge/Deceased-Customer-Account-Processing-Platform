#!/usr/bin/env python3
"""
Create a Word document with all persons for easy form filling
"""

import sqlite3
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_reference_docx():
    base_dir = Path(__file__).parent

    # Create document
    doc = Document()

    # Set up styles
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    # Title
    title = doc.add_heading('BeneBridge Mock Database - Complete Person Reference', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    subtitle.add_run('For Creating Death Certificates, IDs, and Beneficiary Forms').bold = True
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # Summary info
    summary = doc.add_paragraph()
    summary.add_run('Total Cases: 20 | Deceased Persons: 20 | Beneficiaries: 43\n').bold = True
    summary.add_run('Institution: Community National Bank | Account Type: IRA Only\n')
    summary.add_run('Fraud Cases: Case 13 (Blacklisted), Case 19 (SSN Mismatch)')

    doc.add_page_break()

    # ========== DECEASED PERSONS ==========
    doc.add_heading('DECEASED PERSONS (All from Los Angeles County, CA)', 1)
    doc.add_paragraph('Use this section to fill out DEATH CERTIFICATES')

    # Get deceased from database
    conn = sqlite3.connect(base_dir / "dmf_mock.db")
    conn.row_factory = sqlite3.Row
    c = conn.cursor()
    c.execute("SELECT * FROM deceased_persons ORDER BY case_id")
    deceased_list = [dict(row) for row in c.fetchall()]
    conn.close()

    # Get accounts for each deceased
    conn = sqlite3.connect(base_dir / "financial_accounts.db")
    conn.row_factory = sqlite3.Row
    c = conn.cursor()

    for person in deceased_list:
        # Case header
        heading = doc.add_heading(f'Case #{person["case_id"]}: {person["full_name"]}', 2)
        heading.runs[0].font.color.rgb = RGBColor(0, 0, 128)

        # Create table for deceased info
        table = doc.add_table(rows=12, cols=2)
        table.style = 'Light Grid Accent 1'

        rows = table.rows

        # Fill table
        rows[0].cells[0].text = 'Full Name'
        rows[0].cells[1].text = person['full_name']

        rows[1].cells[0].text = 'First Name'
        rows[1].cells[1].text = person['first_name']

        rows[2].cells[0].text = 'Last Name'
        rows[2].cells[1].text = person['last_name']

        rows[3].cells[0].text = 'Gender'
        rows[3].cells[1].text = 'Male' if person['gender'] == 'M' else 'Female'

        rows[4].cells[0].text = 'SSN'
        rows[4].cells[1].text = person['ssn']

        rows[5].cells[0].text = 'Date of Birth'
        rows[5].cells[1].text = person['dob']

        rows[6].cells[0].text = 'Date of Death'
        rows[6].cells[1].text = person['dod']

        rows[7].cells[0].text = 'Address'
        rows[7].cells[1].text = f"{person['address_street']}, {person['address_city']}, {person['address_state']} {person['address_zip']}"

        rows[8].cells[0].text = 'City/County'
        rows[8].cells[1].text = f"{person['address_city']}, Los Angeles County, CA"

        rows[9].cells[0].text = 'Height'
        rows[9].cells[1].text = person['height']

        rows[10].cells[0].text = 'Weight'
        rows[10].cells[1].text = f"{person['weight']} lbs"

        rows[11].cells[0].text = 'Eye Color'
        rows[11].cells[1].text = person['eye_color']

        # Get accounts for this deceased
        c.execute("SELECT * FROM accounts WHERE deceased_case_id = ? ORDER BY balance DESC", (person['case_id'],))
        accounts = [dict(row) for row in c.fetchall()]

        if accounts:
            doc.add_paragraph()
            doc.add_paragraph('Financial Accounts:').bold = True
            for acc in accounts:
                acc_para = doc.add_paragraph(style='List Bullet')
                acc_para.add_run(f"{acc['account_id']}: ")
                acc_para.add_run(f"Community National Bank IRA - ${acc['balance']:,.2f}")

        doc.add_paragraph()
        doc.add_paragraph('_' * 80)
        doc.add_paragraph()

    conn.close()

    doc.add_page_break()

    # ========== BENEFICIARIES ==========
    doc.add_heading('BENEFICIARIES', 1)
    doc.add_paragraph('Use this section to fill out ID DOCUMENTS and BENEFICIARY CLAIM FORMS')

    # Get beneficiaries
    conn = sqlite3.connect(base_dir / "beneficiary_registry.db")
    conn.row_factory = sqlite3.Row
    c = conn.cursor()
    c.execute("SELECT * FROM beneficiaries ORDER BY beneficiary_id")
    beneficiaries_list = [dict(row) for row in c.fetchall()]
    conn.close()

    # Get identity records
    conn = sqlite3.connect(base_dir / "identity_verification.db")
    conn.row_factory = sqlite3.Row
    c = conn.cursor()
    c.execute("SELECT * FROM identity_records ORDER BY id")
    identities = [dict(row) for row in c.fetchall()]
    conn.close()

    # Create identity lookup by SSN
    identity_map = {i['ssn']: i for i in identities}

    for person in beneficiaries_list:
        # Beneficiary header
        heading = doc.add_heading(f'Beneficiary #{person["beneficiary_id"]}: {person["full_name"]}', 2)
        heading.runs[0].font.color.rgb = RGBColor(0, 100, 0)

        # Get identity record
        identity = identity_map.get(person['ssn'])

        # Create table
        table = doc.add_table(rows=16, cols=2)
        table.style = 'Light Grid Accent 1'

        rows = table.rows

        # Fill table
        rows[0].cells[0].text = 'Full Name'
        rows[0].cells[1].text = person['full_name']

        rows[1].cells[0].text = 'First Name'
        rows[1].cells[1].text = person['first_name']

        rows[2].cells[0].text = 'Last Name'
        rows[2].cells[1].text = person['last_name']

        rows[3].cells[0].text = 'Gender'
        rows[3].cells[1].text = 'Male' if person['gender'] == 'M' else 'Female'

        rows[4].cells[0].text = 'SSN'
        rows[4].cells[1].text = person['ssn']

        rows[5].cells[0].text = 'Date of Birth'
        rows[5].cells[1].text = person['dob']

        rows[6].cells[0].text = 'Address'
        rows[6].cells[1].text = f"{person['address_street']}, {person['address_city']}, {person['address_state']} {person['address_zip']}"

        rows[7].cells[0].text = 'Phone'
        rows[7].cells[1].text = person['phone']

        rows[8].cells[0].text = 'Email'
        rows[8].cells[1].text = person['email']

        if identity:
            rows[9].cells[0].text = 'Driver\'s License #'
            rows[9].cells[1].text = identity['drivers_license']

            rows[10].cells[0].text = 'DL State'
            rows[10].cells[1].text = identity['dl_state']

            rows[11].cells[0].text = 'DL Expiration'
            rows[11].cells[1].text = identity['dl_expiration']

            rows[12].cells[0].text = 'Height'
            rows[12].cells[1].text = identity['height']

            rows[13].cells[0].text = 'Weight'
            rows[13].cells[1].text = f"{identity['weight']} lbs"

            rows[14].cells[0].text = 'Eye Color'
            rows[14].cells[1].text = identity['eye_color']

            rows[15].cells[0].text = 'Photo Reference'
            rows[15].cells[1].text = identity['photo_reference']

        doc.add_paragraph()
        doc.add_paragraph('_' * 80)
        doc.add_paragraph()

    # Save document
    output_file = base_dir / "PERSON_REFERENCE_FOR_FORMS.docx"
    doc.save(output_file)

    print("=" * 80)
    print("WORD DOCUMENT CREATED")
    print("=" * 80)
    print(f"File: {output_file}")
    print(f"Size: {output_file.stat().st_size / 1024:.1f} KB")
    print()
    print("Document contains:")
    print(f"  - {len(deceased_list)} deceased persons (for death certificates)")
    print(f"  - {len(beneficiaries_list)} beneficiaries (for IDs and claim forms)")
    print()
    print("Ready to use for form filling!")

if __name__ == "__main__":
    create_reference_docx()
