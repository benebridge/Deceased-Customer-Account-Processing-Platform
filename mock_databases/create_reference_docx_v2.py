#!/usr/bin/env python3
"""
Create comprehensive Word document with all persons for form filling
Includes all fields needed for CA driver's licenses and beneficiary claim forms
"""

import sqlite3
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_reference_docx():
    base_dir = Path(__file__).parent

    # Create document
    doc = Document()

    # Set up styles
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    # Title
    title = doc.add_heading('BeneBridge Mock Database - Complete Person Reference', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    subtitle.add_run('For Creating Death Certificates, CA Driver\'s Licenses, and IRA Beneficiary Claim Forms').bold = True
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # Summary info
    summary = doc.add_paragraph()
    summary.add_run('Total Cases: 20 | Deceased Persons: 20 | Beneficiaries: 47\n').bold = True
    summary.add_run('All persons from California | Institution: Community National Bank | Account Type: IRA Only\n')
    summary.add_run('Fraud Cases: Case 13, Case 19')

    doc.add_page_break()

    # ========== DECEASED PERSONS ==========
    doc.add_heading('SECTION 1: DECEASED PERSONS', 1)
    doc.add_paragraph('Use for LA COUNTY DEATH CERTIFICATES')
    doc.add_paragraph('All from Los Angeles County, CA')
    doc.add_paragraph()

    # Get deceased from database
    conn = sqlite3.connect(base_dir / "dmf_mock.db")
    conn.row_factory = sqlite3.Row
    c = conn.cursor()
    c.execute("SELECT * FROM deceased_persons ORDER BY case_id")
    deceased_list = [dict(row) for row in c.fetchall()]
    conn.close()

    # Get accounts for each deceased
    conn_fin = sqlite3.connect(base_dir / "financial_accounts.db")
    conn_fin.row_factory = sqlite3.Row
    c_fin = conn_fin.cursor()

    for person in deceased_list:
        # Case header
        heading = doc.add_heading(f'Case #{person["case_id"]}: {person["full_name"]}', 2)
        heading.runs[0].font.color.rgb = RGBColor(0, 0, 128)

        # Create table for deceased info
        table = doc.add_table(rows=13, cols=2)
        table.style = 'Light Grid Accent 1'
        rows = table.rows

        # Fill table
        rows[0].cells[0].text = 'Full Name'
        rows[0].cells[1].text = person['full_name']

        rows[1].cells[0].text = 'First Name'
        rows[1].cells[1].text = person['first_name']

        rows[2].cells[0].text = 'Last Name'
        rows[2].cells[1].text = person['last_name']

        rows[3].cells[0].text = 'Gender'
        rows[3].cells[1].text = 'Male' if person['gender'] == 'M' else 'Female'

        rows[4].cells[0].text = 'SSN'
        rows[4].cells[1].text = person['ssn']

        rows[5].cells[0].text = 'Date of Birth'
        rows[5].cells[1].text = person['dob']

        rows[6].cells[0].text = 'Date of Death'
        rows[6].cells[1].text = person['dod']

        rows[7].cells[0].text = 'Address'
        rows[7].cells[1].text = f"{person['address_street']}, {person['address_city']}, {person['address_state']} {person['address_zip']}"

        rows[8].cells[0].text = 'City/County'
        rows[8].cells[1].text = f"{person['address_city']}, Los Angeles County, CA"

        rows[9].cells[0].text = 'Height'
        rows[9].cells[1].text = person['height']

        rows[10].cells[0].text = 'Weight'
        rows[10].cells[1].text = f"{person['weight']} lbs"

        rows[11].cells[0].text = 'Eye Color'
        rows[11].cells[1].text = person['eye_color']

        rows[12].cells[0].text = 'Sex'
        rows[12].cells[1].text = person['gender']

        # Get accounts for this deceased
        c_fin.execute("SELECT * FROM accounts WHERE deceased_case_id = ? ORDER BY balance DESC", (person['case_id'],))
        accounts = [dict(row) for row in c_fin.fetchall()]

        if accounts:
            doc.add_paragraph()
            doc.add_paragraph('IRA Accounts:').bold = True
            for acc in accounts:
                acc_para = doc.add_paragraph(style='List Bullet')
                acc_para.add_run(f"Account #: {acc['account_id']} | ")
                acc_para.add_run(f"Balance: ${acc['balance']:,.2f}")

        doc.add_paragraph()
        doc.add_paragraph('_' * 80)
        doc.add_paragraph()

    conn_fin.close()

    doc.add_page_break()

    # ========== BENEFICIARIES ==========
    doc.add_heading('SECTION 2: BENEFICIARIES', 1)
    doc.add_paragraph('Use for CA DRIVER\'S LICENSES and IRA BENEFICIARY CLAIM FORMS')
    doc.add_paragraph('All from California')
    doc.add_paragraph()

    # Get beneficiaries
    conn_ben = sqlite3.connect(base_dir / "beneficiary_registry.db")
    conn_ben.row_factory = sqlite3.Row
    c_ben = conn_ben.cursor()
    c_ben.execute("SELECT * FROM beneficiaries ORDER BY beneficiary_id")
    beneficiaries_list = [dict(row) for row in c_ben.fetchall()]

    # Get beneficiary designations
    c_ben.execute("SELECT * FROM beneficiary_designations")
    designations = [dict(row) for row in c_ben.fetchall()]
    conn_ben.close()

    # Get identity records
    conn_id = sqlite3.connect(base_dir / "identity_verification.db")
    conn_id.row_factory = sqlite3.Row
    c_id = conn_id.cursor()
    c_id.execute("SELECT * FROM identity_records ORDER BY id")
    identities = [dict(row) for row in c_id.fetchall()]
    conn_id.close()

    # Create identity lookup by SSN
    identity_map = {i['ssn']: i for i in identities}

    for person in beneficiaries_list:
        # Beneficiary header
        heading = doc.add_heading(f'Beneficiary #{person["beneficiary_id"]}: {person["full_name"]}', 2)
        heading.runs[0].font.color.rgb = RGBColor(0, 100, 0)

        # Get identity record
        identity = identity_map.get(person['ssn'])

        # Get which deceased case this beneficiary is linked to
        deceased_case = person['deceased_case_id']
        deceased_person = next((d for d in deceased_list if d['case_id'] == deceased_case), None)

        # Get account designations for this beneficiary
        ben_accounts = [d for d in designations if d['beneficiary_id'] == person['beneficiary_id']]

        # Header info
        doc.add_paragraph(f"Relationship to Case #{deceased_case} ({deceased_person['full_name'] if deceased_person else 'Unknown'}): {person['relationship']}").italic = True

        doc.add_paragraph()
        doc.add_paragraph('FOR CA DRIVER\'S LICENSE:').bold = True

        # Driver's License table
        table_dl = doc.add_table(rows=13, cols=2)
        table_dl.style = 'Light Grid Accent 1'
        rows_dl = table_dl.rows

        rows_dl[0].cells[0].text = 'Driver\'s License Number'
        rows_dl[0].cells[1].text = identity['drivers_license'] if identity else ''

        rows_dl[1].cells[0].text = 'Issue Date'
        rows_dl[1].cells[1].text = identity['dl_issue_date'] if identity else ''

        rows_dl[2].cells[0].text = 'Expiration Date'
        rows_dl[2].cells[1].text = identity['dl_expiration'] if identity else ''

        rows_dl[3].cells[0].text = 'Last Name'
        rows_dl[3].cells[1].text = person['last_name']

        rows_dl[4].cells[0].text = 'First Name'
        rows_dl[4].cells[1].text = person['first_name']

        rows_dl[5].cells[0].text = 'Address'
        rows_dl[5].cells[1].text = f"{person['address_street']}, {person['address_city']}, {person['address_state']} {person['address_zip']}"

        rows_dl[6].cells[0].text = 'Date of Birth'
        rows_dl[6].cells[1].text = person['dob']

        rows_dl[7].cells[0].text = 'Sex'
        rows_dl[7].cells[1].text = person['gender']

        rows_dl[8].cells[0].text = 'Height'
        rows_dl[8].cells[1].text = identity['height'] if identity else ''

        rows_dl[9].cells[0].text = 'Weight'
        rows_dl[9].cells[1].text = f"{identity['weight']} lbs" if identity else ''

        rows_dl[10].cells[0].text = 'Eye Color'
        rows_dl[10].cells[1].text = identity['eye_color'] if identity else ''

        rows_dl[11].cells[0].text = 'Hair Color'
        rows_dl[11].cells[1].text = identity['hair_color'] if identity else ''

        rows_dl[12].cells[0].text = 'Photo'
        rows_dl[12].cells[1].text = identity['photo_reference'] if identity else ''

        doc.add_paragraph()
        doc.add_paragraph('FOR IRA BENEFICIARY CLAIM FORM:').bold = True

        # Beneficiary claim form table
        table_claim = doc.add_table(rows=7, cols=2)
        table_claim.style = 'Light Grid Accent 1'
        rows_claim = table_claim.rows

        rows_claim[0].cells[0].text = 'Full Name'
        rows_claim[0].cells[1].text = person['full_name']

        rows_claim[1].cells[0].text = 'SSN'
        rows_claim[1].cells[1].text = person['ssn']

        rows_claim[2].cells[0].text = 'Email'
        rows_claim[2].cells[1].text = person['email']

        rows_claim[3].cells[0].text = 'Phone'
        rows_claim[3].cells[1].text = person['phone']

        rows_claim[4].cells[0].text = 'Relationship to Deceased'
        rows_claim[4].cells[1].text = person['relationship']

        # List accounts
        if ben_accounts:
            account_list = []
            for acc in ben_accounts:
                account_list.append(f"{acc['account_id']} ({acc['percentage']}%)")
            rows_claim[5].cells[0].text = 'IRA Account Number(s)'
            rows_claim[5].cells[1].text = ', '.join(account_list)

            rows_claim[6].cells[0].text = 'Beneficiary Percentage(s)'
            percentages = ', '.join([f"{acc['percentage']}%" for acc in ben_accounts])
            rows_claim[6].cells[1].text = percentages

        doc.add_paragraph()
        doc.add_paragraph('_' * 80)
        doc.add_paragraph()

    # Save document
    output_file = base_dir / "PERSON_REFERENCE_FOR_FORMS.docx"
    doc.save(output_file)

    print("=" * 80)
    print("WORD DOCUMENT CREATED")
    print("=" * 80)
    print(f"File: {output_file}")
    print(f"Size: {output_file.stat().st_size / 1024:.1f} KB")
    print()
    print("Document contains:")
    print(f"  - {len(deceased_list)} deceased persons (for LA County death certificates)")
    print(f"  - {len(beneficiaries_list)} beneficiaries (for CA driver's licenses and IRA claim forms)")
    print()
    print("SECTION 1: Deceased persons with:")
    print("  - Full name, first, last, gender, SSN, DOB, DOD")
    print("  - Address, city/county, height, weight, eye color")
    print("  - IRA account numbers and balances")
    print()
    print("SECTION 2: Beneficiaries with:")
    print("  - FOR CA DRIVER'S LICENSE:")
    print("    DL number, issue date, expiration, name, address, DOB")
    print("    Sex, height, weight, eye color, hair color, photo")
    print("  - FOR IRA BENEFICIARY CLAIM FORM:")
    print("    Full name, SSN, email, phone, relationship")
    print("    Account numbers, beneficiary percentages")
    print()
    print("All beneficiaries are now from California!")
    print("Ready to use for form filling!")

if __name__ == "__main__":
    create_reference_docx()
